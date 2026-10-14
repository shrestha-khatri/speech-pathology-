"""
Young Healthy Control Numerical Analysis Script

This script performs statistical analysis on the extracted speech features
for young healthy control participants.

The analysis includes:
• Descriptive summary statistics by task
• Mann–Whitney U tests (simple vs reading)
• Cohen’s d effect size calculation
• Multi-sheet Excel report generation
• Basic Excel formatting for significance highlighting

Input:
youngHC_paper_features.csv

Output:
YoungHC_Numerical_Analysis_<timestamp>.xlsx containing:
    Summary statistics
    Statistical test results
    Full descriptive statistics
    Interpretation table
    Raw dataset
"""

import pandas as pd
import numpy as np
from scipy.stats import mannwhitneyu
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime


print("Loading CSV and computing stats")
df = pd.read_csv(r"C:\Users\DELL\Codes\DH604\young healthy analysis\youngHC_paper_features.csv")


# Summary statistics grouped by task
print("Computing summary statistics")
numerical = df.groupby('task')[['spir', 'ppr', 'hnr', 'jitter_local', 'shimmer_local']].agg(
    ['mean', 'std', 'count', 'min', 'max', 'median']
).round(3)


# Mann–Whitney U statistical tests (simple vs reading)
print("Running Mann-Whitney U tests")
p_values = []

for feat in ['spir', 'ppr', 'hnr', 'jitter_local', 'shimmer_local']:
    simple = df[df.task == 'simple'][feat].dropna()
    reading = df[df.task == 'reading'][feat].dropna()
    
    if len(simple) >= 3 and len(reading) >= 3:
        stat, p = mannwhitneyu(simple, reading)
        
        n1, n2 = len(simple), len(reading)
        pooled_std = np.sqrt(((n1-1)*simple.std()**2 + (n2-1)*reading.std()**2) / (n1+n2-2))
        cohens_d = (simple.mean() - reading.mean()) / pooled_std if pooled_std > 0 else 0
        
        p_values.append({
            'Feature': feat.upper(),
            'Simple_Mean': round(simple.mean(), 3),
            'Reading_Mean': round(reading.mean(), 3),
            'Difference': round(simple.mean() - reading.mean(), 3),
            'Simple_n': n1,
            'Reading_n': n2,
            'U_Statistic': round(stat, 1),
            'p_value': round(p, 4),
            'Significant': 'YES' if p < 0.05 else 'NO',
            'Cohens_d': round(cohens_d, 3),
            'Effect_Size': 'LARGE' if abs(cohens_d) > 0.8 else ('MEDIUM' if abs(cohens_d) > 0.5 else 'SMALL')
        })
    else:
        p_values.append({
            'Feature': feat.upper(),
            'Simple_Mean': 0,
            'Reading_Mean': 0,
            'Difference': 0,
            'Simple_n': len(simple),
            'Reading_n': len(reading),
            'U_Statistic': 0,
            'p_value': 1.0,
            'Significant': 'NO DATA',
            'Cohens_d': 0,
            'Effect_Size': 'N/A'
        })

p_values_df = pd.DataFrame(p_values)


# Full descriptive statistics
print("Creating task comparison")
task_comparison = df.groupby('task')[['spir', 'ppr', 'hnr', 'jitter_local', 'shimmer_local']].describe().round(3)


# Interpretation table summarising means
interp_data = {
    'Feature': ['SPIR', 'PPR', 'HNR', 'Jitter_Local', 'Shimmer_Local'],
    'Definition': [
        'Speech Pauses/min (Kovac)',
        'Pause Proportion Ratio',
        'Harmonics-to-Noise (Favaro)',
        'Local cycle irregularity (Favaro)',
        'Local amplitude variation (Favaro)'
    ],
    'Simple_Mean': [
        df[df.task=='simple']['spir'].mean(),
        df[df.task=='simple']['ppr'].mean(),
        df[df.task=='simple']['hnr'].mean(),
        df[df.task=='simple']['jitter_local'].mean(),
        df[df.task=='simple']['shimmer_local'].mean()
    ],
    'Reading_Mean': [
        df[df.task=='reading']['spir'].mean(),
        df[df.task=='reading']['ppr'].mean(),
        df[df.task=='reading']['hnr'].mean(),
        df[df.task=='reading']['jitter_local'].mean(),
        df[df.task=='reading']['shimmer_local'].mean()
    ],
    'Expected_PD': ['↑ Higher', '↑ Higher', '↓ Lower', '↑ Higher', '↑ Higher'],
    'Paper': ['Kovac', 'Kovac', 'Favaro', 'Favaro', 'Favaro']
}
interp_df = pd.DataFrame(interp_data).round(3)


# Write Excel workbook
print("Writing Excel")
excel_file = f'YoungHC_Numerical_Analysis_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx'

with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
    numerical.to_excel(writer, sheet_name='Summary_Stats')
    p_values_df.to_excel(writer, sheet_name='P_Values_EffectSize', index=False)
    task_comparison.to_excel(writer, sheet_name='Full_Descriptive')
    interp_df.to_excel(writer, sheet_name='Interpretation', index=False)
    df.to_excel(writer, sheet_name='Raw_Data', index=False)

print(f"Excel saved: {excel_file}")


# Apply basic Excel formatting
print("Applying basic formatting")
wb = load_workbook(excel_file)

header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
sig_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
nonsig_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

ws = wb['P_Values_EffectSize']
for col in range(1, min(12, ws.max_column + 1)):
    ws.cell(row=1, column=col).fill = header_fill
    ws.cell(row=1, column=col).font = header_font

for row in range(2, ws.max_row + 1):
    cell = ws.cell(row=row, column=9)
    val = str(cell.value).upper()
    if 'YES' in val:
        cell.fill = sig_fill
    elif 'NO' in val:
        cell.fill = nonsig_fill

wb.save(excel_file)


# Console preview
print("\nSignificant Features:")
sig_results = p_values_df[p_values_df['Significant'] == 'YES']

if len(sig_results) > 0:
    for _, row in sig_results.iterrows():
        print(f"{row['Feature']}: simple={row['Simple_Mean']:.1f}, reading={row['Reading_Mean']:.1f}, p={row['p_value']:.4f}")
else:
    print("No significant differences detected")

print(f"\nFull report: {excel_file}")
